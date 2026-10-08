from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"


def set_font(run, name="Aptos", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=8.8):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    set_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)


def add_table(doc, headers, rows, widths=None, fill="E6EDF7"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, fill)
        set_cell_text(cell, header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.8)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, len(row) - 1) and len(row) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[idx], value, align=align)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


def add_title(doc, title, subtitle):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    set_font(run, name="Aptos Display", size=19, bold=True)
    p.paragraph_format.space_after = Pt(3)
    sub = doc.add_paragraph()
    r = sub.add_run(subtitle)
    set_font(r, size=10.5, color="555555")
    sub.paragraph_format.space_after = Pt(9)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_font(run, name="Aptos Display" if level == 1 else "Aptos", size=14 if level == 1 else 11.5, bold=True)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)


def para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.06
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, size=10.2)


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_font(run, size=9.8)


def setup():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10.2)
    return doc


def create_doc():
    doc = setup()

    add_title(
        doc,
        "Tester Findings and QA Visibility Report",
        "Presentation document for UI, UX, workflow, and K6 load testing observations",
    )
    para(doc, "This document is prepared from the tester perspective for tomorrow's presentation. It records issues that became visible only while using the University Management System end to end across admin, faculty, and student roles. Several findings were not obvious during development or ordinary QA checks because the screens individually appeared complete, while the defects emerged only when the full user journey was exercised.")
    heading(doc, "Testing Context")
    bullets(doc, [
        "Admin flows were tested for master setup, subject allocation, timetable, examination, evaluation setup, and reports.",
        "Faculty flows were tested for allocated subjects, timetable, students, attendance, assignments, study materials, marks, and feedback summaries.",
        "Student flows were tested for dashboard, timetable, attendance, assignments, study materials, calendar, results, and feedback submission.",
        "The main testing principle was simple: every dropdown should come from a master or allocation source, and every user should see only related data.",
    ])
    heading(doc, "Finding Categories")
    add_table(
        doc,
        ["Category", "Purpose", "Presentation Focus"],
        [
            ["Suggested UI changes", "Improve visible alignment, action clarity, and screen polish", "What the user can immediately see"],
            ["Suggested UX changes", "Improve guidance, validation, feedback, and confidence", "What the user understands while using the app"],
            ["Workflow changes", "Fix broken or incomplete end-to-end flows", "What fails only across role-to-role journeys"],
            ["Concurrency and load testing", "Validate behavior under repeated and simultaneous actions", "What happens when many users act together"],
        ],
        widths=[1.55, 2.6, 2.25],
        fill="D9EAF7",
    )

    doc.add_page_break()
    heading(doc, "Suggested UI Changes")
    para(doc, "The UI issues identified by the tester were mostly not logic failures. They were presentation and interaction problems that can reduce confidence during a live demo. A developer may miss them because the data technically loads, and QA may miss them if the screen is checked only for functional availability.")
    add_table(
        doc,
        ["Screen or Area", "Tester Observation", "Suggested UI Change", "Priority"],
        [
            ["Study materials", "Resource card button looked clickable but initially had no action", "Ensure every visible button has an action, disabled state, or message", "High"],
            ["Forms and modals", "Validation sometimes appeared outside the active popup context", "Show validation inside the modal near the Save button or affected field", "High"],
            ["Dashboard cards", "Some cards showed static or outdated values during role switching", "Use live counts and remove hardcoded schedule or metric values", "High"],
            ["Tables", "Long text could make table rows look uneven", "Use consistent column widths and keep action columns compact", "Medium"],
            ["Navigation", "Some module names were similar, such as resources and study materials", "Use role-specific labels while keeping a common layout", "Medium"],
        ],
        widths=[1.35, 2.15, 2.45, 0.65],
        fill="D9EAF7",
    )
    heading(doc, "UI Acceptance Points")
    bullets(doc, [
        "Every primary button should visibly complete an action or explain why the action cannot run.",
        "Cards should not contain placeholder filenames or unrelated test labels during the final presentation.",
        "Tables should remain readable at the presentation screen resolution without horizontal confusion.",
        "Status chips should use consistent labels such as Active, Pending, Submitted, Completed, and In Progress.",
    ])

    doc.add_page_break()
    heading(doc, "Suggested UX Changes")
    para(doc, "The UX findings came from observing how a real user would complete a task rather than checking whether a page exists. The tester found that the app needs stronger feedback after actions, clearer empty states, and better guidance when users move between related modules.")
    add_table(
        doc,
        ["UX Issue", "Why It Matters", "Suggested Change"],
        [
            ["Unclear success feedback", "Users may not know whether attendance, feedback, or upload was saved", "Show concise success messages and refresh affected data immediately"],
            ["Weak empty states", "A blank table can look like a broken page", "Explain whether no data exists, data is scoped away, or the user must select a filter"],
            ["Role transition confusion", "Admin, faculty, and student screens share similar modules with different powers", "Use role-specific descriptions and scoped dropdown labels"],
            ["First-time form use", "Users may not understand dependency between subject, batch, and roster", "Use step indicators for flows such as attendance and assignment creation"],
            ["Error language", "Technical errors can be hard to present", "Convert validation and authorization failures into readable messages"],
        ],
        widths=[1.65, 2.35, 2.5],
        fill="E8EEF9",
    )
    heading(doc, "Tester UX Notes")
    bullets(doc, [
        "The user should never need to inspect the browser console to understand a failure.",
        "A successful save should update the list or card on the same screen without forcing manual refresh.",
        "Filter-driven screens should clearly show which subject, batch, or professor scope is currently active.",
        "Student-facing pages should avoid administrative words when the student only needs to view or submit.",
    ])

    doc.add_page_break()
    heading(doc, "Workflow Changes and Errors Found During Flow Testing")
    para(doc, "Workflow testing exposed problems that were not visible when pages were checked independently. The main lesson is that the system must be tested as a chain: admin configures data, faculty acts on that data, and students see the result.")
    add_table(
        doc,
        ["Flow", "Error or Gap Found", "Workflow Change Suggested", "Impact"],
        [
            ["Attendance", "Initial behavior could fail when saving the same slot again", "Allow attendance update for an already submitted slot", "Prevents demo interruption"],
            ["Assignments", "Faculty could create work but reviewing submissions needed clearer flow", "Add submission review panel from assignment card", "Completes faculty loop"],
            ["Resources", "Student resource visibility depended on correct subject and batch relation", "Scope resources by allocation and batch before showing to students", "Protects relevance"],
            ["Feedback", "Students should not review unrelated professors", "Generate feedback options only from assigned subject faculty allocations", "Protects validity"],
            ["Calendar", "Calendar should not depend on static examples", "Feed it from timetable, assignments, exams, components, and faculty events", "Improves student value"],
        ],
        widths=[1.15, 2.2, 2.25, 0.9],
        fill="D9EAF7",
    )
    heading(doc, "Why These Were Not Immediately Visible")
    bullets(doc, [
        "A developer sees that components compile and screens render, but not always whether the cross-role chain is complete.",
        "QA may verify forms individually, but hidden defects appear when the next role depends on the previous role's output.",
        "Several issues were state-based: they appeared only after creating, updating, refreshing, or logging in as another user.",
        "The most important defects were not visual crashes; they were broken continuity between modules.",
    ])

    doc.add_page_break()
    heading(doc, "Concurrency and Load Testing Issues Using Grafana K6")
    para(doc, "K6 load testing should be used to simulate repeated logins, reads, submissions, and upload-related actions. The tester should present this as a risk-focused check rather than a full production benchmark. The purpose is to identify whether the current presentation baseline remains stable under simultaneous use.")
    add_table(
        doc,
        ["K6 Scenario", "Risk Being Tested", "Potential Issue", "Suggested Threshold"],
        [
            ["Concurrent login", "Many students or faculty sign in before demo", "Auth delay or token errors", "95 percent under 800 ms"],
            ["Attendance submit", "Multiple faculty submit class records at the same time", "Duplicate or locked records", "No failed saves for valid requests"],
            ["Student calendar load", "Students open calendar after faculty updates", "Slow event aggregation", "95 percent under 1.2 s"],
            ["Assignment submission", "Many students upload assignment files", "Upload timeout or file handling failure", "No failed uploads under file limit"],
            ["Resource download", "Students download study material together", "Slow response or missing file errors", "No 404 for valid resources"],
        ],
        widths=[1.35, 1.95, 1.85, 1.35],
        fill="E8EEF9",
    )
    heading(doc, "K6 Reporting Format")
    bullets(doc, [
        "Report virtual users, duration, total requests, failed request rate, average response time, and p95 response time.",
        "Separate read-heavy tests from write-heavy tests because attendance and submissions can reveal different problems.",
        "Flag any endpoint that returns authorization errors for valid scoped users.",
        "Use Grafana dashboards to show response time trend, error rate, request rate, and slowest endpoint.",
    ])
    para(doc, "Final conclusion: The tester's contribution is to expose the gap between pages that appear ready and workflows that are actually reliable. The most important improvements are clickable actions, visible validation, role-scoped continuity, update-safe attendance, complete assignment submission review, and basic K6 evidence for concurrency confidence.")

    return doc


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    path = OUT / "Tester_QA_Findings_Presentation_Document.docx"
    create_doc().save(path)
    print(path)
