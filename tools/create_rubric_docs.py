from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"


def set_font(run, name="Aptos", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    set_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)


def style_table(table, header_fill="D9EAF7"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_border(cell)
            if i == 0:
                shade_cell(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_font(run, name="Aptos Display" if level == 1 else "Aptos", size=15 if level == 1 else 12, bold=True)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.space_after = Pt(7)
    if bold_lead:
        lead = p.add_run(bold_lead)
        set_font(lead, bold=True)
        text = text.removeprefix(bold_lead).lstrip()
    run = p.add_run(text)
    set_font(run)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_font(run, size=10)


def add_small_table(doc, headers, rows, widths=None, fill="D9EAF7"):
    table = doc.add_table(rows=1, cols=len(headers))
    style_table(table, fill)
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.2)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, len(row) - 1) and len(row) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[idx], value, align=align, size=9.0)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


def add_title(doc, title, subtitle):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    set_font(run, name="Aptos Display", size=20, bold=True)
    p.paragraph_format.space_after = Pt(4)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(10)
    r = sub.add_run(subtitle)
    set_font(r, size=11, color="555555")


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    return doc


def page_break(doc):
    doc.add_page_break()


def tester_doc():
    doc = setup_doc()
    add_title(doc, "Tester Rubric Document", "Five page testing plan aligned to the 5 mark tester evaluation rubric")
    add_para(doc, "This document explains how testing will be planned, executed, recorded, and reported for the University Management System. It is organized directly against the tester rubric so each marking criterion has visible evidence and a clear method of assessment.")
    add_heading(doc, "Rubric Alignment Summary")
    add_small_table(
        doc,
        ["Criterion", "Marks", "Evidence to Provide"],
        [
            ["Test case design and coverage", "1.5", "Scenario matrix covering login, masters, attendance, calendar, assignments, resources, and feedback"],
            ["Test inputs expected and actual outputs", "1", "Structured records showing input values, expected behavior, actual result, and pass or fail status"],
            ["Edge boundary and exceptional cases", "0.5", "Negative tests for missing fields, duplicate values, invalid dates, unauthorized access, and file issues"],
            ["Performance and validation testing", "1", "Basic response time checks, form validation checks, and file upload size checks"],
            ["Ability to identify testing issues and explain results", "1", "Defect log with severity, cause, affected flow, and recommended fix"],
        ],
        widths=[3.1, 0.7, 3.0],
        fill="F8CB9C",
    )
    add_heading(doc, "Testing Scope")
    add_bullets(doc, [
        "Authentication and role routing for admin, faculty, and student users.",
        "Master-driven dropdowns for departments, designations, programs, batches, courses, subjects, and academic sessions.",
        "Role-scoped faculty views for subjects, batches, students, timetable, attendance, assignments, marks, resources, and feedback.",
        "Student views for attendance, calendar events, assignments, resources, results, and professor feedback.",
    ])

    page_break(doc)
    add_heading(doc, "Test Case Design and Coverage")
    add_para(doc, "The test suite should be organized by user journey rather than only by screen. This keeps the evidence close to how the application is actually demonstrated: admin creates or configures data, faculty acts on that data, and students view or respond to it.")
    add_small_table(
        doc,
        ["Flow", "Primary Tests", "Coverage Goal"],
        [
            ["Login", "Valid login, invalid password, first login password change, forgot password screen", "Confirm correct routing and visible validation"],
            ["Masters", "Create, edit, delete, duplicate record validation, dropdown population", "Ensure all form values come from master data"],
            ["Attendance", "Slot selection, roster load, submit, update existing attendance, student summary", "Prove admin and faculty actions reach student view"],
            ["Calendar", "Timetable classes, exams, assignments, quizzes, faculty events", "Show one consolidated student academic calendar"],
            ["Assignments", "Faculty creates assignment, student uploads, faculty views submissions", "Prove coursework loop is complete"],
            ["Resources", "Faculty uploads resource, student downloads relevant resource", "Prove subject and batch scoped sharing"],
            ["Feedback", "Student reviews own professor, faculty sees anonymous summary", "Prove feedback cannot target unrelated professors"],
        ],
        widths=[1.25, 3.7, 1.8],
        fill="F8CB9C",
    )
    add_para(doc, "Coverage is considered adequate when every role has at least one successful path and one controlled failure path for the features assigned to that role.")

    page_break(doc)
    add_heading(doc, "Test Inputs Expected Outputs and Actual Outputs")
    add_para(doc, "Every executed test should capture exact inputs. For form-based tests, this includes selected dropdown values, typed field values, uploaded file names, and the active role. For API-backed flows, the visible UI result is the primary evidence, supported by network response messages when required.")
    add_small_table(
        doc,
        ["ID", "Input", "Expected Output", "Actual Output", "Result"],
        [
            ["T01", "Faculty login AVV-F1024 with valid password", "Faculty dashboard opens with scoped data", "Dashboard loads", "Pass"],
            ["T02", "Create student with existing email", "Popup shows duplicate email message", "Validation message shown", "Pass"],
            ["T03", "Mark attendance for CSE 2022 Section A", "Attendance saved and summary updates", "Attendance saved", "Pass"],
            ["T04", "Student opens resources after faculty upload", "Relevant resource card is visible", "Resource visible and downloadable", "Pass"],
            ["T05", "Student submits feedback for assigned professor", "Anonymous feedback saved", "Feedback saved", "Pass"],
        ],
        widths=[0.55, 2.0, 2.0, 1.55, 0.7],
        fill="F8CB9C",
    )
    add_heading(doc, "Recording Standard")
    add_bullets(doc, [
        "Use Pass only when the visible result matches the expected result.",
        "Use Fail when the result is wrong, missing, blocked, or visible only after refresh when it should be immediate.",
        "Record Not Run only when a test was not attempted, never when the tester is unsure of the result.",
    ])

    page_break(doc)
    add_heading(doc, "Edge Boundary and Exceptional Cases")
    add_para(doc, "Boundary testing verifies that the system fails cleanly. For this application, the most important exceptional cases are invalid role access, incomplete forms, duplicate master values, missing file uploads, invalid dates, oversized files, and access to unrelated faculty or student data.")
    add_small_table(
        doc,
        ["Area", "Boundary or Exception", "Expected Handling"],
        [
            ["Authentication", "Wrong password or unknown user", "Login stays on page and shows an error message"],
            ["Masters", "Duplicate employee ID, student ID, subject code, or email", "Modal shows validation message without background-only errors"],
            ["Attendance", "Faculty selects subject not allocated to them", "Request is blocked by authorization"],
            ["Assignments", "Student attempts unrelated batch assignment", "Submission is forbidden"],
            ["Resources", "Resource has no uploaded file", "Button shows a clear message rather than doing nothing"],
            ["Calendar", "End date before start date", "Form rejects the event with a readable message"],
        ],
        widths=[1.3, 2.6, 2.7],
        fill="F8CB9C",
    )
    add_heading(doc, "Performance and Validation Testing")
    add_bullets(doc, [
        "Dashboard, list, and form pages should load without visible delay during the demo dataset.",
        "Attendance roster loading should remain smooth for the assigned class size.",
        "File upload should respect the configured maximum file size and return a readable error when exceeded.",
        "Dropdowns should populate from backend master or scoped allocation endpoints rather than static screen data.",
    ])

    page_break(doc)
    add_heading(doc, "Issue Identification and Result Explanation")
    add_para(doc, "The tester should explain not only whether a case passed, but also what the result means for the user journey. A strong issue report identifies the role, screen, action, observed behavior, expected behavior, likely impact, and recommended correction.")
    add_small_table(
        doc,
        ["Severity", "Meaning", "Example"],
        [
            ["High", "Blocks a presentation-critical flow", "Faculty cannot submit attendance or student cannot view assignment"],
            ["Medium", "Feature works but result is incomplete or confusing", "Validation appears behind popup or resource button gives no feedback"],
            ["Low", "Cosmetic or minor usability issue", "Text spacing, label alignment, or unclear button wording"],
        ],
        widths=[1.0, 2.35, 3.15],
        fill="F8CB9C",
    )
    add_heading(doc, "Final Submission Checklist")
    add_bullets(doc, [
        "Attach the completed test case table with expected and actual outputs.",
        "Attach evidence screenshots for at least one pass and one controlled failure per major flow.",
        "Confirm that master dropdowns and role-scoped data are tested.",
        "Summarize unresolved issues with severity and recommended next action.",
    ])
    add_para(doc, "Conclusion: The tester earns full marks when the evidence shows broad coverage, exact input and output comparison, meaningful edge testing, basic performance validation, and clear explanation of defects and results.")
    return doc


def qa_doc():
    doc = setup_doc()
    add_title(doc, "QA Rubric Document", "Five page quality assurance plan aligned to the 5 mark QA evaluation rubric")
    add_para(doc, "This document defines the QA approach for validating the University Management System before demonstration. It focuses on data reliability, output correctness, error identification, performance confidence, and practical recommendations for improvement.")
    add_heading(doc, "Rubric Alignment Summary")
    add_small_table(
        doc,
        ["Criterion", "Marks", "Evidence to Provide"],
        [
            ["Data quality and preprocessing validation", "1", "Checks for completeness, consistency, uniqueness, referential integrity, and master-data accuracy"],
            ["Model or output quality verification", "1", "Verification that UI outputs and reports match database state and role expectations"],
            ["Identification of errors limitations and inconsistencies", "1", "Logged defects, limitations, edge cases, and unresolved risks"],
            ["Reliability and performance checks", "1", "Repeated-flow checks, reload checks, response confidence, and file handling checks"],
            ["Documentation and recommendations for improvement", "1", "Clear QA report with findings, evidence, and practical improvement plan"],
        ],
        widths=[3.1, 0.7, 3.0],
        fill="BDD0EE",
    )
    add_heading(doc, "QA Objective")
    add_para(doc, "The QA objective is to confirm that the system behaves consistently across admin, faculty, and student roles. The application should not merely display screens; it should preserve correct relationships between masters, allocations, timetable slots, attendance records, assignments, resources, and feedback.")
    add_bullets(doc, [
        "Master data must be the source for dropdowns and repeated values.",
        "Faculty must see only subjects, batches, students, and activities related to their allocations.",
        "Students must see only their own batch, attendance, coursework, calendar items, resources, results, and feedback options.",
    ])

    page_break(doc)
    add_heading(doc, "Data Quality and Preprocessing Validation")
    add_para(doc, "Data quality validation checks whether the records used by the application are fit for use. In this system, the most important data relationships are between users, staff, students, batches, subjects, subject allocations, timetable entries, assignments, and submissions.")
    add_small_table(
        doc,
        ["Data Area", "QA Check", "Acceptance Rule"],
        [
            ["Users", "Every active staff or student login maps to one profile", "No orphaned profile is used for demo login"],
            ["Masters", "Department, designation, program, academic year, batch, course, and subject values are unique", "Dropdowns contain clean non-duplicate values"],
            ["Allocations", "Subject allocation connects subject, faculty, batch, academic year, and semester", "Faculty scope derives from allocation records"],
            ["Attendance", "Attendance record uses valid student, subject, batch, date, and session", "No unrelated student appears in roster"],
            ["Assignments", "Assignment uses allocated subject and batch", "Students see assignments only for their batch"],
            ["Feedback", "Student feedback target is an assigned professor", "Feedback options exclude unrelated faculty"],
        ],
        widths=[1.2, 3.1, 2.2],
        fill="BDD0EE",
    )
    add_para(doc, "Preprocessing for demo data should include removing accidental duplicates, resetting known demo credentials, confirming current batch semester values, and ensuring test records do not pollute final reports.")

    page_break(doc)
    add_heading(doc, "Model and Output Quality Verification")
    add_para(doc, "Output verification compares what the user sees with the intended backend state. A screen passes QA only when visible counts, lists, dropdowns, buttons, and status labels match the data that should be available to the current role.")
    add_small_table(
        doc,
        ["Output", "Verification Method", "Expected Quality"],
        [
            ["Faculty dashboard", "Compare assigned subjects, classes, students, tasks, and schedule against allocations", "Counts and schedule are role scoped"],
            ["Student calendar", "Compare classes, exams, assignments, assessment components, and events", "Calendar is descriptive and batch specific"],
            ["Attendance summary", "Compare present, absent, and OD records with percentage display", "Percentages are accurate and update after submit"],
            ["Assignment cards", "Compare submissions with progress and status", "Student and faculty views show correct status"],
            ["Resources", "Upload as faculty and view as student", "Only relevant resources are visible and download works"],
            ["Feedback summary", "Submit as student and view as faculty", "Faculty sees aggregate only, not student identity"],
        ],
        widths=[1.45, 3.2, 1.85],
        fill="BDD0EE",
    )
    add_heading(doc, "Visual and Interaction Quality")
    add_bullets(doc, [
        "Buttons should produce a visible action or a clear message.",
        "Validation errors should appear in the active popup or active form area.",
        "Empty states should explain why no records are visible.",
        "Tables and cards should remain aligned at desktop presentation resolution.",
    ])

    page_break(doc)
    add_heading(doc, "Errors Limitations and Inconsistencies")
    add_para(doc, "QA must identify errors that affect correctness and limitations that may affect confidence. The report should separate blocking failures from acceptable presentation limitations so that the team can prioritize fixes.")
    add_small_table(
        doc,
        ["Category", "Example", "QA Response"],
        [
            ["Error", "A faculty member sees subjects allocated to another faculty member", "Mark high severity and block demo path until fixed"],
            ["Inconsistency", "A dropdown value is typed statically on one page but master-driven elsewhere", "Mark medium severity and align source"],
            ["Limitation", "Reports export is basic and not fully polished", "Document limitation and keep out of main demo unless needed"],
            ["Usability gap", "A button has no action or no feedback", "Add click behavior or clear disabled state"],
            ["Data issue", "A temporary test record appears in final report", "Clean data or label test records before demo"],
        ],
        widths=[1.2, 3.0, 2.3],
        fill="BDD0EE",
    )
    add_heading(doc, "Reliability and Performance Checks")
    add_bullets(doc, [
        "Repeat the same attendance submit path twice to confirm updates do not fail on duplicate records.",
        "Refresh after resource upload, assignment submission, feedback submission, and marks save to confirm persistence.",
        "Check role switching by logging out and signing in as admin, faculty, and student.",
        "Confirm upload and download work for at least one small document file.",
    ])

    page_break(doc)
    add_heading(doc, "Documentation and Recommendations")
    add_para(doc, "The final QA document should be readable as a decision record. It should state what was checked, what passed, what remains risky, and what should be improved after the presentation baseline.")
    add_small_table(
        doc,
        ["Recommendation", "Reason", "Priority"],
        [
            ["Keep all dropdowns master or allocation driven", "Prevents inconsistent data entry and strengthens role logic", "High"],
            ["Maintain a small demo credential sheet", "Avoids failed logins during presentation", "High"],
            ["Add export polish after core flows", "Reports are useful but not critical to the current demo story", "Medium"],
            ["Add richer empty loading and error states", "Improves confidence during live navigation", "Medium"],
            ["Add audit logs for submissions and feedback", "Useful for future production reliability", "Low"],
        ],
        widths=[2.4, 3.2, 0.9],
        fill="BDD0EE",
    )
    add_heading(doc, "Final QA Sign Off Checklist")
    add_bullets(doc, [
        "All seven presentation flows have at least one verified pass path.",
        "Role-scoped data is verified for one faculty account and one student account.",
        "Known limitations are documented with priority and suggested next action.",
        "Screens used in the demo are free from obvious broken buttons, background-only errors, and layout misalignment.",
    ])
    add_para(doc, "Conclusion: QA earns full marks when the evidence proves reliable data, correct outputs, clear defect identification, repeatable behavior, and useful recommendations for improvement.")
    return doc


def save(doc, filename):
    path = OUT / filename
    doc.save(path)
    return path


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    print(save(tester_doc(), "Tester_Rubric_Document.docx"))
    print(save(qa_doc(), "QA_Rubric_Document.docx"))
