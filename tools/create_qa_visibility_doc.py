from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"


def set_font(run, name="Aptos", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=8.8):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    set_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)


def add_table(doc, headers, rows, widths=None, fill="D9EAF7"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, fill)
        set_cell_text(cell, header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.7)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, len(row) - 1) and len(row) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[idx], value, align=align, size=8.6)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


def setup():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10.2)
    return doc


def add_title(doc, title, subtitle):
    p = doc.add_paragraph(style="Title")
    run = p.add_run(title)
    set_font(run, name="Aptos Display", size=19, bold=True)
    p.paragraph_format.space_after = Pt(3)
    sub = doc.add_paragraph()
    r = sub.add_run(subtitle)
    set_font(r, size=10.5, color="555555")
    sub.paragraph_format.space_after = Pt(9)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_font(run, name="Aptos Display" if level == 1 else "Aptos", size=14 if level == 1 else 11.5, bold=True)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)


def para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.06
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, size=10.2)


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_font(run, size=9.8)


def create_doc():
    doc = setup()

    add_title(
        doc,
        "QA Validation and Release Readiness Report",
        "Presentation document for quality checks, missed visibility, workflow risks, and K6 validation",
    )
    para(
        doc,
        "This document is prepared from the QA perspective for tomorrow's presentation. It explains what QA validated after the tester surfaced end-to-end issues, which quality risks were not obvious from screen-level checks, and what must be verified before the University Management System can be presented as a stable role-based application.",
    )
    heading(doc, "QA Scope")
    bullets(
        doc,
        [
            "QA reviewed admin, faculty, and student flows as connected journeys rather than isolated pages.",
            "QA checked whether master data, subject allocation, timetable, attendance, assignments, resources, feedback, and calendar data stayed consistent across roles.",
            "QA verified that role scoping was respected: faculty should see only related classes and students, and students should see only their own academic data.",
            "QA treated any hidden browser-console error, unclear validation message, or stale dashboard value as a release-readiness concern.",
        ],
    )
    heading(doc, "QA Evidence Categories")
    add_table(
        doc,
        ["Category", "QA Purpose", "Presentation Focus"],
        [
            ["Suggested UI changes", "Confirm that screens are readable, aligned, and visibly actionable", "Quality of what appears on screen"],
            ["Suggested UX changes", "Confirm that users understand each action, result, and error", "Quality of user confidence"],
            ["Workflow changes", "Confirm that one role's action correctly affects the next role", "Quality of end-to-end continuity"],
            ["Concurrency and load testing", "Confirm that critical flows stay stable under repeated use", "Quality under simultaneous activity"],
        ],
        widths=[1.6, 2.65, 2.15],
        fill="D9EAF7",
    )

    doc.add_page_break()
    heading(doc, "Suggested UI Changes")
    para(
        doc,
        "QA found that several issues were not visible as backend failures. They appeared as presentation quality problems that could reduce confidence during a demo even when the data was technically available.",
    )
    add_table(
        doc,
        ["Screen or Area", "QA Observation", "Suggested UI Change", "Severity"],
        [
            ["Login and role entry", "Common login should not force users to choose role manually", "Infer the user role after authentication and route automatically", "High"],
            ["Dashboards", "Cards could show outdated or placeholder counts", "Bind cards to current scoped data and remove static examples", "High"],
            ["Forms and modals", "Validation could appear outside the popup or behind the active form", "Keep all validation, loading, and success states inside the active modal", "High"],
            ["Tables", "Action columns and long names were visually inconsistent", "Use fixed action width, readable wrapping, and compact status chips", "Medium"],
            ["Navigation", "Shared modules needed consistent naming across roles", "Use common layout with role-specific labels and active states", "Medium"],
        ],
        widths=[1.35, 2.15, 2.45, 0.65],
        fill="E8EEF9",
    )
    heading(doc, "QA UI Acceptance Criteria")
    bullets(
        doc,
        [
            "No primary action should appear clickable unless it performs an action or explains why it is disabled.",
            "No page should show test labels, incorrect names, stale dates, or placeholder dashboard values.",
            "Every modal should contain its own validation message and should not leave errors behind on the background page.",
            "The same visual language should be used for status, empty states, filters, and actions across all roles.",
        ],
    )

    doc.add_page_break()
    heading(doc, "Suggested UX Changes")
    para(
        doc,
        "QA focused on whether a real user can complete the task without help from a developer. The main concern was not that pages were missing, but that some actions needed clearer confirmation, better empty states, and stronger guidance between dependent steps.",
    )
    add_table(
        doc,
        ["UX Risk", "QA Finding", "Recommended Change"],
        [
            ["Unclear completion", "Users may not know whether attendance, upload, feedback, or assignment submission saved", "Show success messages and refresh the affected list or card immediately"],
            ["Weak empty states", "Blank tables can look broken when data is simply scoped away", "Explain whether the user needs to select a subject, batch, slot, or date"],
            ["Dependency confusion", "Subject, batch, semester, faculty, and student rosters depend on prior master data", "Use dependent dropdown labels and disable later fields until earlier choices are valid"],
            ["First login flow", "Temporary password change can display the wrong name if session data is stale", "Use authenticated user data only and avoid incorrect names"],
            ["Error handling", "Technical API errors are difficult to present", "Convert validation and authorization errors into readable modal messages"],
        ],
        widths=[1.55, 2.35, 2.6],
        fill="D9EAF7",
    )
    heading(doc, "QA UX Acceptance Criteria")
    bullets(
        doc,
        [
            "A user should never need to inspect the browser console or network tab to understand what failed.",
            "A successful action should visibly change the screen state within the same workflow.",
            "A blocked action should clearly explain the missing prerequisite, such as subject allocation or batch selection.",
            "Role-specific screens should use language that matches what the user can actually do.",
        ],
    )

    doc.add_page_break()
    heading(doc, "Workflow Changes and Errors Found During QA Validation")
    para(
        doc,
        "QA validation treated the system as a chain of responsibility. Admin creates the structure, faculty uses the structure, and students consume the result. The most important defects were found where that chain crossed from one role to another.",
    )
    add_table(
        doc,
        ["Flow", "QA Issue Found", "Workflow Change Required", "Release Risk"],
        [
            ["Attendance", "Slots, batches, faculty subjects, and student rosters must remain linked", "Admin assigns slots; faculty marks only allocated batches; students see only their own attendance", "High"],
            ["Student calendar", "Calendar must be fed by assignments, timetable, exams, quizzes, and deadlines", "Create faculty/admin event inputs and aggregate events by student batch", "High"],
            ["Evaluation pattern", "Components need consistent admin and faculty setup", "Define components per subject or batch and expose marks entry only to assigned faculty", "High"],
            ["Resources", "Students should not see unrelated subject materials", "Scope uploads by subject allocation, batch, and semester", "Medium"],
            ["Anonymous feedback", "Students should only review their own professors", "Generate feedback forms from actual faculty allocation at semester end", "Medium"],
        ],
        widths=[1.15, 2.2, 2.25, 0.9],
        fill="E8EEF9",
    )
    heading(doc, "Why QA Could Miss It Earlier")
    bullets(
        doc,
        [
            "Individual screens can pass a basic check while the full admin-to-faculty-to-student chain still fails.",
            "Several defects require switching users after creating data, which is not always part of a page-level QA pass.",
            "Role scope problems appear only when more than one faculty, batch, subject, or student account exists.",
            "The biggest risk is not a page crash; it is showing valid-looking but unrelated data to the wrong user.",
        ],
    )

    doc.add_page_break()
    heading(doc, "Concurrency and Load Testing Issues Using Grafana K6")
    para(
        doc,
        "QA should use Grafana K6 to validate whether the presentation baseline can survive repeated use by multiple users. The goal is not a production capacity certificate. The goal is to catch preventable failures in login, attendance, calendar, assignment upload, and resource download before the demo.",
    )
    add_table(
        doc,
        ["K6 Scenario", "QA Risk", "Possible Failure", "Acceptance Target"],
        [
            ["Concurrent login", "Many users sign in near the same time", "Slow auth response or token failure", "p95 under 800 ms with low error rate"],
            ["Attendance submit", "Multiple faculty submit or update attendance", "Duplicate rows, locks, or failed writes", "No failed valid save and no duplicate slot record"],
            ["Calendar load", "Students load aggregated academic events", "Slow event aggregation or missing events", "p95 under 1.2 s for scoped users"],
            ["Assignment upload", "Many students submit files close to deadline", "Timeout, file failure, or bad status", "No failed uploads within file limit"],
            ["Resource download", "Students download materials together", "404 for valid files or delayed response", "No missing valid resource and stable p95"],
        ],
        widths=[1.35, 1.9, 1.9, 1.35],
        fill="D9EAF7",
    )
    heading(doc, "QA Reporting Format")
    bullets(
        doc,
        [
            "Report virtual users, test duration, total requests, failed request rate, p95 response time, and slowest endpoint.",
            "Separate read-heavy tests from write-heavy tests so dashboard and calendar issues do not hide attendance and upload failures.",
            "Flag every authorization failure for a valid user as a role-scope defect, not just a performance defect.",
            "Use Grafana panels for response time trend, request rate, error rate, and endpoint breakdown.",
        ],
    )
    para(
        doc,
        "Final conclusion: QA's contribution is to convert tester observations into release-readiness checks. The core recommendation is to validate every critical workflow through role-scoped data, clear UI feedback, readable error handling, and basic K6 evidence before presenting the system.",
    )
    return doc


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    output_path = OUT / "QA_Validation_Release_Readiness_Report.docx"
    create_doc().save(output_path)
    print(output_path)
