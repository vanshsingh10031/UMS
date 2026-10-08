from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


OUT_DIR = Path("outputs/demo_dataset")
OUT_DIR.mkdir(parents=True, exist_ok=True)

NOW = datetime(2026, 10, 9, 0, 0, 0)
TEMP_PASSWORD = "amrita2026"


def row(**kwargs):
    base = {"CreatedAt": NOW, "UpdatedAt": NOW}
    base.update(kwargs)
    return base


def initials(name: str) -> str:
    return "".join(part[0] for part in name.replace("Dr. ", "").split()[:2]).upper()


def sql_value(value):
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, datetime):
        return f"'{value.strftime('%Y-%m-%dT%H:%M:%S')}'"
    if isinstance(value, date):
        return f"'{value.isoformat()}'"
    if isinstance(value, time):
        return f"'{value.strftime('%H:%M:%S')}'"
    if isinstance(value, (int, float)):
        return str(value)
    escaped = str(value).replace("'", "''")
    return f"N'{escaped}'"


def insert_block(table, rows):
    if not rows:
        return []
    cols = list(rows[0].keys())
    lines = [f"SET IDENTITY_INSERT dbo.{table} ON;"]
    for item in rows:
        values = ", ".join(sql_value(item.get(col)) for col in cols)
        lines.append(f"INSERT INTO dbo.{table} ({', '.join(cols)}) VALUES ({values});")
    lines.append(f"SET IDENTITY_INSERT dbo.{table} OFF;")
    lines.append("")
    return lines


tables = defaultdict(list)

roles = [(1, "Admin", "University administrator"), (2, "Faculty", "Faculty member"), (3, "Student", "Student user")]
for rid, name, desc in roles:
    tables["RoleMasters"].append(row(Id=rid, Name=name, Description=desc))

departments = [
    (1, "AIDS", "Artificial Intelligence and Data Science"),
    (2, "AIDSM", "AI and Data Science Medical"),
    (3, "CSE", "Computer Science"),
    (4, "MAT", "Mathematics"),
    (5, "MGT", "Management"),
]
for did, code, name in departments:
    tables["DepartmentMasters"].append(row(Id=did, Code=code, Name=name))

for idx, name in enumerate(["Professor", "Associate Professor", "Assistant Professor", "Lab Instructor", "Administrative Staff"], 1):
    tables["DesignationMasters"].append(row(Id=idx, Name=name))

programs = [
    (1, "BT-AIDS", "B.Tech AIDS"),
    (2, "BT-AIDSM", "B.Tech AIDS Medical"),
]
for pid, code, name in programs:
    tables["ProgramMasters"].append(row(Id=pid, Code=code, Name=name))

sessions = [(1, "2025-2029", 1), (2, "2024-2028", 2)]
for sid, ay, sem in sessions:
    tables["AcademicSessions"].append(row(Id=sid, AcademicYear=ay, Semester=sem, IsActive=True))

courses = [
    (1, "AIDS-BT", "B.Tech Artificial Intelligence and Data Science", 160, "Artificial Intelligence and Data Science"),
    (2, "AIDSM-BT", "B.Tech AI and Data Science Medical", 160, "AI and Data Science Medical"),
]
for cid, code, name, credits, dept in courses:
    tables["Courses"].append(row(Id=cid, CourseCode=code, Name=name, Credits=credits, Department=dept, Description=f"{name} undergraduate programme"))

batches = [
    (1, "AIDS 2025 - Batch A", "B.Tech AIDS", "2025-2029", 1, 40),
    (2, "AIDS 2025 - Batch B", "B.Tech AIDS", "2025-2029", 1, 40),
    (3, "AIDS Medical 2025 - Batch A", "B.Tech AIDS Medical", "2025-2029", 1, 40),
    (4, "AIDS 2024 - Batch A", "B.Tech AIDS", "2024-2028", 2, 40),
    (5, "AIDS Medical 2024 - Batch A", "B.Tech AIDS Medical", "2024-2028", 2, 40),
]
for bid, name, program, ay, sem, _size in batches:
    tables["Batches"].append(row(Id=bid, Name=name, Program=program, AcademicYear=ay, Semester=sem))

subjects = [
    (1, "25AID101", "Programming for AI", 4, 1, 1),
    (2, "25AID102", "Linear Algebra for Data Science", 4, 1, 1),
    (3, "25AID103", "Digital Systems", 3, 1, 1),
    (4, "25AID104", "Statistics for Analytics", 4, 1, 1),
    (5, "25AID105", "Data Science Foundations", 4, 1, 1),
    (6, "25AIDM101", "Medical Data Foundations", 4, 1, 2),
    (7, "25AIDM102", "Human Biology for AI", 3, 1, 2),
    (8, "25AIDM103", "Healthcare Systems", 3, 1, 2),
    (9, "24AID201", "Data Structures for AI", 4, 2, 1),
    (10, "24AID202", "Database Systems", 4, 2, 1),
    (11, "24AID203", "Machine Learning Basics", 4, 2, 1),
    (12, "24AID204", "Probability Models", 3, 2, 1),
    (13, "24AID205", "Web Application Engineering", 3, 2, 1),
    (14, "24AIDM201", "Clinical Data Management", 4, 2, 2),
    (15, "24AIDM202", "Biomedical Signal Processing", 4, 2, 2),
    (16, "24AIDM203", "AI for Medical Imaging", 4, 2, 2),
]
for sid, code, name, credits, sem, course_id in subjects:
    tables["Subjects"].append(row(Id=sid, Code=code, Name=name, Credits=credits, Semester=sem, CourseId=course_id))

users = [row(Id=1, Username="admin@amrita.edu", Email="admin@amrita.edu", Role="Admin", RoleMasterId=1, DisplayName="Dr. Meera Nair", PasswordHash=TEMP_PASSWORD, MustChangePassword=False)]
staff_names = [
    "Dr. Arjun Menon", "Dr. Lakshmi Iyer", "Dr. Vivek Raman", "Dr. Priya Nair", "Dr. Karthik Rao",
    "Dr. Sneha Kapoor", "Dr. Rohan Das", "Dr. Mehul Shah", "Dr. Nisha Varma", "Dr. Sanjay Pillai",
    "Dr. Aditi Kulkarni", "Dr. Rahul Bose", "Dr. Farah Khan", "Dr. Neeraj S", "Dr. Kavya Krishnan",
]
designation_cycle = ["Professor", "Associate Professor", "Assistant Professor"]
for i, name in enumerate(staff_names, 1):
    username = f"AVV-F{1200+i:04d}"
    email = f"{name.replace('Dr. ', '').lower().replace(' ', '.')}@amrita.edu"
    users.append(row(Id=i + 1, Username=username, Email=email, Role="Faculty", RoleMasterId=2, DisplayName=name, PasswordHash=TEMP_PASSWORD, MustChangePassword=True))
    dept = "AI and Data Science Medical" if i in {6, 7, 8, 14, 15} else "Artificial Intelligence and Data Science"
    tables["Staff"].append(row(Id=i, EmployeeId=username, Name=name, Email=email, Department=dept, Designation=designation_cycle[(i - 1) % 3], UserId=i + 1))

first_names = ["Ananya", "Ishaan", "Diya", "Kiran", "Naina", "Aarav", "Meera", "Rishi", "Dev", "Tara", "Kabir", "Ira", "Nikhil", "Asha", "Vivaan", "Sana", "Aditya", "Riya", "Varun", "Maya"]
last_names = ["Rao", "Nair", "Menon", "Das", "Singh", "Iyer", "Kapoor", "Pillai", "Shah", "Bose", "Kumar", "George", "Mathew", "Joseph", "Krishnan", "Verma", "Patel", "Reddy", "Khan", "Shetty"]

student_id = 1
for bid, _bname, program, ay, sem, size in batches:
    year_code = "25" if sem == 1 else "24"
    prog_code = "AIDM" if "Medical" in program else "AID"
    dept = "AI and Data Science Medical" if "Medical" in program else "Artificial Intelligence and Data Science"
    for n in range(1, size + 1):
        name = f"{first_names[(student_id - 1) % len(first_names)]} {last_names[(student_id + bid - 2) % len(last_names)]}"
        sid_code = f"CB.EN.U4{prog_code}{year_code}{student_id:03d}"
        email = f"{sid_code.lower()}@dl.students.amrita.edu"
        uid = 16 + student_id
        users.append(row(Id=uid, Username=sid_code, Email=email, Role="Student", RoleMasterId=3, DisplayName=name, PasswordHash=TEMP_PASSWORD, MustChangePassword=True))
        tables["Students"].append(row(Id=student_id, StudentId=sid_code, Name=name, Email=email, Department=dept, Program=program, BatchId=bid, Semester=sem, UserId=uid))
        student_id += 1

tables["Users"].extend(users)

subject_by_id = {s[0]: s for s in subjects}
batch_subjects = {
    1: [1, 2, 3, 4, 5],
    2: [1, 2, 3, 4, 5],
    3: [2, 4, 6, 7, 8],
    4: [9, 10, 11, 12, 13],
    5: [10, 12, 14, 15, 16],
}

alloc_id = 1
subject_staff = {}
for batch_id, subj_ids in batch_subjects.items():
    for index, subject_id in enumerate(subj_ids):
        staff_id = ((subject_id + batch_id + index) % 15) + 1
        subject_staff[(batch_id, subject_id)] = staff_id
        batch = tables["Batches"][batch_id - 1]
        tables["SubjectFacultyAllocations"].append(row(Id=alloc_id, SubjectId=subject_id, StaffId=staff_id, BatchId=batch_id, AcademicYear=batch["AcademicYear"], Semester=batch["Semester"]))
        alloc_id += 1

enroll_id = 1
for student in tables["Students"]:
    for subject_id in batch_subjects[student["BatchId"]]:
        tables["Enrollments"].append(row(Id=enroll_id, StudentId=student["Id"], SubjectId=subject_id))
        enroll_id += 1

days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
slots = [(time(9, 0), time(10, 0)), (time(10, 0), time(11, 0)), (time(11, 15), time(12, 15)), (time(13, 30), time(14, 30)), (time(14, 45), time(15, 45))]
tt_id = 1
for batch_id, subj_ids in batch_subjects.items():
    for idx, subject_id in enumerate(subj_ids):
        st, et = slots[idx]
        tables["TimetableEntries"].append(row(Id=tt_id, Day=days[idx], StartTime=st, EndTime=et, Room=f"A-{200+batch_id}{idx+1}", SubjectId=subject_id, StaffId=subject_staff[(batch_id, subject_id)], BatchId=batch_id))
        tt_id += 1

attendance_id = 1
start = date(2026, 10, 5)
for batch_id, subj_ids in batch_subjects.items():
    students_in_batch = [s for s in tables["Students"] if s["BatchId"] == batch_id]
    for week in range(2):
        for idx, subject_id in enumerate(subj_ids[:4]):
            class_date = start + timedelta(days=idx + week * 7)
            for student in students_in_batch:
                status = "Absent" if (student["Id"] + subject_id + week) % 17 == 0 else "OD" if (student["Id"] + idx) % 29 == 0 else "Present"
                tables["AttendanceRecords"].append(row(Id=attendance_id, StudentId=student["Id"], SubjectId=subject_id, BatchId=batch_id, Date=class_date, Session="Lecture", Status=status))
                attendance_id += 1

assignment_id = 1
for batch_id, subj_ids in batch_subjects.items():
    for j, subject_id in enumerate(subj_ids[:3]):
        due = datetime(2026, 10, 18 + j * 4 + batch_id, 23, 59, 0)
        title = ["Case Study", "Implementation Lab", "Quiz Preparation"][j]
        subject_name = subject_by_id[subject_id][2]
        tables["Assignments"].append(row(Id=assignment_id, Title=f"{subject_name} {title}", Description=f"Complete the {title.lower()} for {subject_name}.\nReference link: https://lms.amrita.edu/{subject_by_id[subject_id][1].lower()}\nAttachment: {subject_by_id[subject_id][1]}_brief.pdf", SubjectId=subject_id, BatchId=batch_id, StaffId=subject_staff[(batch_id, subject_id)], DueDate=due, MaximumMarks=100 if j != 0 else 50, Rubric="Correctness 40%, documentation 20%, viva 20%, presentation 20%"))
        assignment_id += 1

submission_id = 1
for assignment in tables["Assignments"]:
    students_in_batch = [s for s in tables["Students"] if s["BatchId"] == assignment["BatchId"]]
    for student in students_in_batch[:26]:
        if (student["Id"] + assignment["Id"]) % 5 == 0:
            continue
        tables["Submissions"].append(row(Id=submission_id, AssignmentId=assignment["Id"], StudentId=student["Id"], FileName=f"{student['StudentId']}_{assignment['Id']}.pdf", FilePath=f"uploads/submissions/{student['StudentId']}_{assignment['Id']}.pdf", Notes="Submitted through portal", SubmittedAt=assignment["DueDate"] - timedelta(days=(student["Id"] % 3) + 1)))
        submission_id += 1

mark_id = 1
for enrollment in tables["Enrollments"]:
    student_id_ref = enrollment["StudentId"]
    subject_id_ref = enrollment["SubjectId"]
    max_marks = 100
    score = 58 + ((student_id_ref * 7 + subject_id_ref * 3) % 38)
    status = "Published" if student_id_ref % 4 else "Draft"
    tables["Marks"].append(row(Id=mark_id, StudentId=student_id_ref, SubjectId=subject_id_ref, AssignmentId=None, MarksObtained=score, MaximumMarks=max_marks, Status=status, PublishedAt=NOW if status == "Published" else None))
    mark_id += 1

exam_id = 1
calendar_id = 1
for batch_id, subj_ids in batch_subjects.items():
    for j, subject_id in enumerate(subj_ids[:2]):
        start_dt = datetime(2026, 11, 2 + j + batch_id, 9 + j, 0, 0)
        end_dt = start_dt + timedelta(hours=2)
        tables["Examinations"].append(row(Id=exam_id, Name=f"Internal Assessment {j+1}", SubjectId=subject_id, BatchId=batch_id, StartsAt=start_dt, EndsAt=end_dt, Room=f"EX-{batch_id}{j+1}"))
        exam_id += 1
    for label, typ, offset in [("Class briefing", "Class", 4), ("Quiz window", "Quiz", 8), ("Project deadline", "Deadline", 13)]:
        subject_id = subj_ids[offset % len(subj_ids)]
        starts = datetime(2026, 10, 10 + batch_id + offset, 10, 0, 0)
        tables["AcademicCalendarEvents"].append(row(Id=calendar_id, Title=f"{subject_by_id[subject_id][2]} {label}", Description=f"{typ} item for {tables['Batches'][batch_id-1]['Name']}", Type=typ, Location=f"Room A-{batch_id}{offset}", SubjectId=subject_id, BatchId=batch_id, StaffId=subject_staff[(batch_id, subject_id)], StartsAt=starts, EndsAt=starts + timedelta(hours=1)))
        calendar_id += 1

component_id = 1
for subject in tables["Subjects"]:
    for name, weight, max_marks in [("Internal Assessment 1", 20, 50), ("Assignment", 20, 100), ("End Semester", 60, 100)]:
        tables["EvaluationComponents"].append(row(Id=component_id, Name=name, SubjectId=subject["Id"], Weightage=weight, MaximumMarks=max_marks, Status="Active"))
        component_id += 1

grade_rules = [("A+", 90, 10), ("A", 80, 9), ("B+", 70, 8), ("B", 60, 7), ("C", 50, 6), ("F", 0, 0)]
for i, (grade, min_pct, point) in enumerate(grade_rules, 1):
    tables["GradeRules"].append(row(Id=i, Name=f"Default {grade}", SubjectId=None, MinimumPercentage=min_pct, Grade=grade, GradePoint=point, Status="Active"))

for i, (name, desc, max_marks) in enumerate([("Sports Representation", "Approved sports participation.", 5), ("Medical Consideration", "Approved medical consideration.", 3), ("Hackathon", "Approved technical event participation.", 4)], 1):
    tables["GraceMarkCategories"].append(row(Id=i, Name=name, Description=desc, MaximumMarks=max_marks, Status="Active"))

for i, student in enumerate(tables["Students"][:20], 1):
    subj = batch_subjects[student["BatchId"]][0]
    tables["GraceMarkApplications"].append(row(Id=i, StudentId=student["Id"], SubjectId=subj, GraceMarkCategoryId=(i % 3) + 1, MarksAwarded=(i % 3) + 1, Reason="Approved supporting document submitted", Status="Approved" if i % 4 else "Applied"))

for i in range(1, 21):
    tables["LibraryBooks"].append(row(Id=i, Title=f"Academic Reference Book {i}", Author=f"Author {i}", Isbn=f"UMS-BOOK-{i:04d}", Category="AI and Data Science" if i % 2 else "Medical AI", TotalCount=5 + (i % 4), AvailableCount=3 + (i % 3)))
for i, student in enumerate(tables["Students"][:30], 1):
    tables["LibraryTransactions"].append(row(Id=i, LibraryBookId=(i % 20) + 1, UserId=student["UserId"], IssueDate=NOW - timedelta(days=i % 12), DueDate=NOW + timedelta(days=14 - (i % 5)), ReturnDate=None if i % 3 else NOW - timedelta(days=1), Status="Issued" if i % 3 else "Returned"))

resource_id = 1
for batch_id, subj_ids in batch_subjects.items():
    for subject_id in subj_ids[:3]:
        tables["StudyResources"].append(row(Id=resource_id, Title=f"{subject_by_id[subject_id][2]} lecture notes", Description="Faculty uploaded learning resource.", ResourceType="Notes", SubjectId=subject_id, BatchId=batch_id, StaffId=subject_staff[(batch_id, subject_id)], FileName=f"{subject_by_id[subject_id][1]}_notes.pdf", FilePath=f"uploads/resources/{subject_by_id[subject_id][1]}_notes.pdf"))
        resource_id += 1

feedback_id = 1
comments = ["Clear explanations and useful examples.", "More practice problems would help.", "Good feedback during lab work.", "The course calendar was easy to follow."]
for student in tables["Students"]:
    if student["Id"] % 2:
        subject_id = batch_subjects[student["BatchId"]][0]
        batch = tables["Batches"][student["BatchId"] - 1]
        tables["FacultyFeedback"].append(row(Id=feedback_id, StudentId=student["Id"], StaffId=subject_staff[(student["BatchId"], subject_id)], SubjectId=subject_id, BatchId=student["BatchId"], Rating=3 + (student["Id"] % 3), Comments=comments[student["Id"] % len(comments)], AcademicYear=batch["AcademicYear"], Semester=batch["Semester"], SubmittedAt=NOW - timedelta(days=student["Id"] % 10)))
        feedback_id += 1

survey_id = 1
tables["Surveys"].append(row(Id=survey_id, Title="End Semester Teaching Feedback", Description="Anonymous semester teaching feedback.", Audience="Student", Status="Published", PublishedAt=NOW - timedelta(days=3), ClosesAt=NOW + timedelta(days=30)))
for i, student in enumerate(tables["Students"][:80], 1):
    tables["SurveyResponses"].append(row(Id=i, SurveyId=survey_id, UserId=student["UserId"], ResponseJson='{"rating":4,"comment":"Good course delivery"}', SubmittedAt=NOW - timedelta(days=i % 5)))

for i, user in enumerate(tables["Users"][:80], 1):
    tables["Notifications"].append(row(Id=i, UserId=user["Id"], Title="Demo data notice", Message="Presentation dataset is active for your role-scoped portal.", IsRead=i % 3 == 0))


def excel_value(value):
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    return value


def add_sheet(wb, name, rows):
    ws = wb.create_sheet(name)
    if not rows:
        return ws
    headers = list(rows[0].keys())
    ws.append(headers)
    for item in rows:
        ws.append([excel_value(item.get(h)) for h in headers])
    header_fill = PatternFill("solid", fgColor="8D153A")
    header_font = Font(color="FFFFFF", bold=True, name="Arial", size=10)
    body_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="E6DED9")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, header in enumerate(headers, 1):
        width = min(42, max(12, len(header) + 2, *(len(str(ws.cell(row=r, column=col_idx).value or "")) + 2 for r in range(2, min(ws.max_row, 60) + 1))))
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    return ws


summary_rows = [
    {"Area": "Students", "Count": len(tables["Students"]), "Notes": "200 students across 5 batches"},
    {"Area": "Faculty", "Count": len(tables["Staff"]), "Notes": "15 professors divided across first and second year"},
    {"Area": "Batches", "Count": len(tables["Batches"]), "Notes": "3 first-year batches, 2 second-year batches"},
    {"Area": "Subjects", "Count": len(tables["Subjects"]), "Notes": "8 first-year and 8 second-year subjects"},
    {"Area": "Attendance records", "Count": len(tables["AttendanceRecords"]), "Notes": "Slot-wise attendance sample"},
    {"Area": "Assignments", "Count": len(tables["Assignments"]), "Notes": "Faculty-issued assignments with submissions"},
    {"Area": "Calendar events", "Count": len(tables["AcademicCalendarEvents"]), "Notes": "Classes, quizzes and deadlines"},
    {"Area": "Feedback", "Count": len(tables["FacultyFeedback"]), "Notes": "Students scoped to own professors"},
]

wb = Workbook()
wb.remove(wb.active)
add_sheet(wb, "Summary", summary_rows)
sheet_order = [
    "RoleMasters", "DepartmentMasters", "DesignationMasters", "ProgramMasters", "AcademicSessions",
    "Users", "Staff", "Students", "Courses", "Batches", "Subjects", "SubjectFacultyAllocations",
    "Enrollments", "TimetableEntries", "AttendanceRecords", "Assignments", "Submissions", "Marks",
    "AcademicCalendarEvents", "Examinations", "EvaluationComponents", "GradeRules", "StudyResources",
    "FacultyFeedback", "LibraryBooks", "LibraryTransactions", "Notifications", "GraceMarkCategories",
    "GraceMarkApplications", "Surveys", "SurveyResponses",
]
for table_name in sheet_order:
    add_sheet(wb, table_name[:31], tables[table_name])

workbook_path = OUT_DIR / "UMS_Presentation_Demo_Dataset.xlsx"
wb.save(workbook_path)

insert_order = [
    "RoleMasters", "DepartmentMasters", "DesignationMasters", "ProgramMasters", "AcademicSessions", "Users",
    "Staff", "Courses", "Batches", "Students", "Subjects", "SubjectFacultyAllocations", "Enrollments",
    "TimetableEntries", "AttendanceRecords", "Assignments", "Submissions", "Marks", "Examinations",
    "LibraryBooks", "LibraryTransactions", "Notifications", "AcademicCalendarEvents", "StudyResources",
    "FacultyFeedback", "EvaluationComponents", "GradeRules", "GraceMarkCategories", "GraceMarkApplications",
    "Surveys", "SurveyResponses",
]

sql_lines = [
    "-- UMS presentation replica dataset",
    "-- Target schema: current EF Core UMS.Api migrations",
    "-- Creates a clean database copy named UMSPresentationReplica and fills master, people and workflow data.",
    "USE master;",
    "IF DB_ID(N'UMSPresentationReplica') IS NOT NULL BEGIN ALTER DATABASE UMSPresentationReplica SET SINGLE_USER WITH ROLLBACK IMMEDIATE; DROP DATABASE UMSPresentationReplica; END;",
    "CREATE DATABASE UMSPresentationReplica;",
    "GO",
    "USE UMSPresentationReplica;",
    "GO",
    "-- Run the application's EF migrations against UMSPresentationReplica first, then run the inserts below.",
    "-- Example: dotnet ef database update --project Backend/UMS.Api --connection \"Server=(localdb)\\MSSQLLocalDB;Database=UMSPresentationReplica;Trusted_Connection=True;TrustServerCertificate=True\"",
    "GO",
]
for table_name in insert_order:
    sql_lines.extend(insert_block(table_name, tables[table_name]))

sql_path = OUT_DIR / "UMS_Presentation_Demo_Dataset.sql"
sql_path.write_text("\n".join(sql_lines), encoding="utf-8")

print(f"students={len(tables['Students'])}")
print(f"staff={len(tables['Staff'])}")
print(f"subjects={len(tables['Subjects'])}")
print(f"assignments={len(tables['Assignments'])}")
print(f"attendance_records={len(tables['AttendanceRecords'])}")
print(workbook_path)
print(sql_path)
